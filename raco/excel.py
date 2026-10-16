"""Write every result table to one Excel workbook."""
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

METRICS = ["energy_MWh", "energy_per_year_MWh", "network_lifetime_years", "life_10pct_years", "median_node_life_years",
           "first_failure_obs_years", "n_failed_by_L", "n_failed_by_H", "sla_violation_pct", "overflow_pct",
           "wake_loss_pct", "aging_gini", "aging_std_mV", "mean_active_frac", "wakeups_per_node_day",
           "mean_theta_active", "max_theta", "carbon_op_t", "carbon_emb_t", "carbon_total_t", "ms_per_frame",
           "E_base_MWh", "E_leak_MWh", "E_dyn_MWh", "E_sleep_MWh", "E_wake_MWh"]

SHEET_HELP = [
    ("README", "This sheet"),
    ("Config", "All model / scenario parameters of the base configuration"),
    ("Main_Summary", "E1: mean and std over seeds for each method (Fig. 4, 3, 15)"),
    ("Main_PerSeed", "E1: raw per-seed results"),
    ("Stats_Tests", "Paired t-test and Wilcoxon signed-rank test RACO vs each baseline (needs >= 3 seeds)"),
    ("Ablation_Summary", "E2: RACO and its ablations (Fig. 7)"),
    ("Ablation_PerSeed", "E2: raw per-seed results"),
    ("Pareto_Ltarget", "E3: L_target sweep (Fig. 5)"),
    ("Sens_V", "E4: Lyapunov V sweep (Fig. 8)"),
    ("Sens_TauR", "E5: recovery time-constant (Fig. 9)"),
    ("Sens_Climate", "E6: city climate (Fig. 10)"),
    ("Sens_Load", "E7: load level (Fig. 11)"),
    ("Sens_Compensation", "E8: fixed-V silicon vs AVS (Fig. 16)"),
    ("Sens_nExp", "E9: NBTI time exponent (Fig. 18)"),
    ("Sens_Delta", "E9b: capacity head-room over the forecast, SLA vs energy (Fig. 19)"),
    ("Sens_Lref", "E9c: reference lifetime of the aging calibration (always-on life at T_ref)"),
    ("Baseline_Safety", "E12: AgingUtil with the safety rule (AgingUtil-S) vs AgingUtil and RACO"),
    ("Scalability", "E10: decision time vs N (Fig. 14)"),
    ("SolverGap", "E11: per-frame optimality gap vs exhaustive search (Fig. 13)"),
    ("Nodes_Main", "Per-node results, seed 0, main methods (Fig. 12)"),
    ("Daily_<method>", "Daily network-level time series, seed 0"),
    ("Aging_<method>", "Daily ΔVth [mV] of every node, seed 0 (Fig. 2)"),
    ("Week_Schedule", "Active/sleep matrix of the sample week (Fig. 6)"),
]

METRIC_HELP = {
    "energy_MWh": "Total energy over the simulated horizon",
    "network_lifetime_years": "Time to first node failure (observed, or projected from last-year aging rate)",
    "life_10pct_years": "Time until 10 % of nodes have failed",
    "median_node_life_years": "Median node lifetime",
    "n_failed_by_L": "Nodes failed before the lifetime target",
    "sla_violation_pct": "Requests missing the deadline (overflow to cloud + lost during wake-up)",
    "aging_gini": "Inequality of final ΔVth across nodes (0 = perfectly even)",
    "carbon_total_t": "Operational + amortised embodied carbon over the horizon",
    "wakeups_per_node_day": "Sleep->active transitions per node per day",
}


def _summary(d, by):
    cols = [m for m in METRICS if m in d.columns]
    g = d.groupby(by)[cols]
    mu, sd = g.mean(), g.std()
    out = pd.DataFrame(index=mu.index)
    for c in cols:
        out[c + "_mean"] = mu[c]
        out[c + "_std"] = sd[c]
    out["n_seeds"] = g.size()
    return out.reset_index()


def _stats(d):
    from scipy import stats
    rows = []
    r = d[d.method == "RACO"].set_index("seed")
    for m in d.method.unique():
        if m == "RACO":
            continue
        b = d[d.method == m].set_index("seed")
        common = r.index.intersection(b.index)
        for met in ["energy_MWh", "network_lifetime_years", "sla_violation_pct", "carbon_total_t"]:
            x, y = r.loc[common, met].values, b.loc[common, met].values
            row = dict(baseline=m, metric=met, n=len(common), RACO_mean=x.mean(), baseline_mean=y.mean(),
                       diff_pct=100 * (x.mean() - y.mean()) / abs(y.mean()) if y.mean() else np.nan)
            if len(common) >= 3 and np.any(x != y):
                row["t_test_p"] = stats.ttest_rel(x, y).pvalue
                try:
                    row["wilcoxon_p"] = stats.wilcoxon(x, y).pvalue
                except ValueError:
                    row["wilcoxon_p"] = np.nan
            rows.append(row)
    return pd.DataFrame(rows)


def write_excel(path, df, base_cfg, nodes, logs, gap, scal):
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        rd = pd.DataFrame(SHEET_HELP, columns=["sheet", "content"])
        mh = pd.DataFrame(list(METRIC_HELP.items()), columns=["metric", "definition"])
        rd.to_excel(xw, sheet_name="README", index=False)
        mh.to_excel(xw, sheet_name="README", index=False, startrow=len(rd) + 2)
        pd.DataFrame(list(base_cfg.items()), columns=["parameter", "value"]).astype(str).to_excel(xw, sheet_name="Config", index=False)

        main = df[df.exp == "main"]
        _summary(main, "method").to_excel(xw, sheet_name="Main_Summary", index=False)
        main.to_excel(xw, sheet_name="Main_PerSeed", index=False)
        _stats(main).to_excel(xw, sheet_name="Stats_Tests", index=False)
        abl = pd.concat([main[main.method == "RACO"], df[df.exp == "ablation"]])
        _summary(abl, "method").to_excel(xw, sheet_name="Ablation_Summary", index=False)
        abl.to_excel(xw, sheet_name="Ablation_PerSeed", index=False)
        sb = df[df.exp == "safety_baseline"]
        if len(sb):
            sb = pd.concat([main[main.method.isin(["AgingUtil", "RACO"])], sb])
            _summary(sb, "method").to_excel(xw, sheet_name="Baseline_Safety", index=False)
        for exp, sh in [("pareto", "Pareto_Ltarget"), ("V", "Sens_V"), ("tau_r", "Sens_TauR"), ("climate", "Sens_Climate"),
                        ("load", "Sens_Load"), ("compensation", "Sens_Compensation"), ("n_exp", "Sens_nExp"), ("delta", "Sens_Delta"),
                        ("lref", "Sens_Lref")]:
            d = df[df.exp == exp].copy()
            if len(d):
                d["value"] = d["value"].astype(str)
                _summary(d, ["method", "param", "value"]).to_excel(xw, sheet_name=sh, index=False)
        if scal:
            pd.DataFrame(scal).to_excel(xw, sheet_name="Scalability", index=False)
        if gap:
            g = pd.DataFrame(gap)
            g.to_excel(xw, sheet_name="SolverGap", index=False)
            s = g[[c for c in g.columns if "gap" in c]].describe().T
            s.to_excel(xw, sheet_name="SolverGap", startcol=len(g.columns) + 2)
        if nodes:
            rows = []
            for m, n in nodes.items():
                for i in range(len(n["type"])):
                    rows.append(dict(method=m, node=i, **{k: (v[i] if hasattr(v, "__len__") else v) for k, v in n.items()}))
            pd.DataFrame(rows).to_excel(xw, sheet_name="Nodes_Main", index=False)
        for m, lg in (logs or {}).items():
            nd = len(lg["E"])
            pd.DataFrame(dict(day=np.arange(nd), year=np.arange(nd) / 365, energy_kWh=lg["E"] / 1000,
                              mean_demand=lg["dem"], mean_overflow=lg["over"], mean_active_nodes=lg["nact"],
                              alive_nodes=lg["alive"], mean_dVth_mV=lg["a"].mean(axis=1) * 1e3,
                              max_dVth_mV=lg["a"].max(axis=1) * 1e3)).to_excel(xw, sheet_name=f"Daily_{m}"[:31], index=False)
            a = pd.DataFrame(lg["a"] * 1e3, columns=[f"n{i}_{t}" for i, t in enumerate(nodes[m]["type"])])
            a.insert(0, "day", np.arange(nd))
            a.to_excel(xw, sheet_name=f"Aging_{m}"[:31], index=False)
        wk = []
        for m, lg in (logs or {}).items():
            w = pd.DataFrame(lg["week_s"].astype(int), columns=[f"n{i}_{t}" for i, t in enumerate(nodes[m]["type"])])
            w.insert(0, "frame", np.arange(len(w))); w.insert(0, "method", m); w["demand"] = lg["week_lam"]
            wk.append(w)
        if wk:
            pd.concat(wk).to_excel(xw, sheet_name="Week_Schedule", index=False)

        # cosmetics
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F3864")
                c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.freeze_panes = "A2"
            for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 60)), 1):
                w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(i)].width = min(max(9, w + 1), 45)
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if isinstance(c.value, float):
                        c.number_format = "0.000"
